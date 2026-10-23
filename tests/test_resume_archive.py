"""Source integrity and upload isolation, with a disposable DB and no model calls."""
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

_test_data = tempfile.TemporaryDirectory(prefix="offeros-archive-tests-")
os.environ["OFFEROS_DATA_DIR"] = _test_data.name
os.environ["LLM_MODE"] = "mock"

from fastapi.testclient import TestClient
from server.main import app
from server.resume_evidence import attach_resume_evidence, rule_sections, validate_sections
from server.agents.resume_agent import analyze_resume
from server.agents.judge_agent import _compact_resume
from server.agents.action_agent import _compact_for_action
from server.files import extract_text_from_file

RESUME = """测试候选人 | 产品经理
邮箱：candidate@example.test

教育经历
示例大学，信息管理，本科，2018—2022

工作经历
示例软件公司，产品经理，2022.07—2026.06
完成 32 场用户访谈，人工工单下降 19%。

项目经历
客户门户 2.0
负责需求梳理与灰度上线。

专业技能
SQL、原型设计
"""


class EvidenceTests(unittest.TestCase):
    def test_complete_demo_has_source_bound_facts_in_every_section(self):
        from server.demo_resume import DEMO_RESUME_PROFILE, DEMO_RESUME_TEXT
        from server.resume_evidence import SECTION_LABELS
        sections = DEMO_RESUME_PROFILE["resume_sections"]
        self.assertEqual(set(sections), set(SECTION_LABELS))
        self.assertGreaterEqual(sum(map(len, sections.values())), 45)
        for key, rows in sections.items():
            self.assertTrue(rows, key)
            for row in rows:
                self.assertIn(row["quote"], DEMO_RESUME_TEXT)
                self.assertIn(row["value"], row["quote"])
        self.assertIn("example.test", DEMO_RESUME_TEXT)
        self.assertIn("不可拨打", DEMO_RESUME_TEXT)
        self.assertEqual(DEMO_RESUME_PROFILE["personality_traits"], [])
        from server.agents.action_agent import bind_local_evidence
        report = bind_local_evidence({}, DEMO_RESUME_PROFILE, {})
        self.assertEqual(len(report["resume_rewrite"]), 3)
        for item in report["resume_rewrite"]:
            self.assertIn(item["original"], DEMO_RESUME_TEXT)
            self.assertEqual(item["rewrite"], item["original"])

    def test_work_header_order_does_not_turn_dates_into_jobs(self):
        for role in ("产品经理", "软件工程师", "科技产品经理"):
            with self.subTest(role=role):
                raw = f"工作经历\n示例软件公司 | 2022.07—2026.06 | {role}"
                rows = rule_sections(raw)["work"]
                fields = {row["label"]: row["value"] for row in rows}
                self.assertEqual(fields["任职时间"], "2022.07—2026.06")
                self.assertEqual(fields["职位"], role)
                self.assertEqual(fields["公司"], "示例软件公司")

    def test_unknown_header_values_keep_neutral_labels(self):
        rows = rule_sections("实习经历\nAcme | 2023—2024 | UX specialist")["internships"]
        self.assertTrue(any(row["label"] == "实习时间" and row["value"] == "2023—2024" for row in rows))
        self.assertTrue(any(row["label"] == "经历信息" and row["value"] == "UX specialist" for row in rows))

    def test_rejects_fabricated_value_and_quote(self):
        sections = {"skills": [
            {"label": "技能", "value": "Java", "quote": "SQL、原型设计"},
            {"label": "技能", "value": "英语流利", "quote": "英语流利"},
            {"label": "技能", "value": "SQL", "quote": "SQL、原型设计"},
        ]}
        rows = validate_sections(sections, RESUME)["skills"]
        self.assertEqual([row["value"] for row in rows], ["SQL"])

    def test_layout_whitespace_still_resolves_to_original(self):
        raw = "教育经历\n示例大学， 信息管理\n本科"
        result = validate_sections({"education": [{"label": "学历", "value": "本科", "quote": "示例大学，信息管理 本科"}]}, raw)
        row = result["education"][0]
        self.assertEqual(row["quote"], "示例大学， 信息管理\n本科")
        self.assertEqual((row["line_start"], row["line_end"]), (2, 3))

    def test_sections_preserve_experience_without_inventing_missing_skills(self):
        result = rule_sections(RESUME)
        self.assertTrue(any("人工工单下降 19%" in row["quote"] for row in result["work"]))
        self.assertIn("客户门户 2.0", result["projects"][0]["quote"])
        self.assertEqual(result["awards"], [])
        self.assertNotIn("英语", str(result))

    def test_untrusted_legacy_profile_does_not_add_resume_facts(self):
        profile = attach_resume_evidence({"skills": ["英语流利"], "evidence": ["编造经历"], "personality_traits": ["外向"]}, RESUME)
        self.assertNotIn("英语流利", str(profile["skills"]))
        self.assertNotIn("编造经历", str(profile["evidence"]))
        self.assertEqual(profile["personality_traits"], [])
        self.assertNotIn("resume_sections", _compact_resume(profile))
        self.assertNotIn("candidate@example.test", str(_compact_resume(profile)))
        self.assertEqual(_compact_for_action(profile, {})[0], _compact_resume(profile))

    def test_local_resume_parsing_never_uses_model_fabrications(self):
        output = {"resume_sections": {"skills": [
            {"label": "语言", "value": "英语流利", "quote": "英语流利"},
            {"label": "技能", "value": "SQL", "quote": "SQL、原型设计"},
        ]}}
        with patch("server.llm.llm.complete_json", return_value=(output, {"used": True})) as complete:
            profile, _ = analyze_resume(RESUME, "test.txt")
        complete.assert_not_called()
        self.assertNotIn("英语流利", str(profile))
        self.assertIn("SQL", str(profile["resume_sections"]["skills"]))
        self.assertFalse(profile["parse_meta"]["used"])

    def test_word_tables_are_extracted_in_document_order(self):
        import docx
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "table.docx"
            document = docx.Document()
            document.add_paragraph("教育经历")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "示例大学"
            table.cell(0, 1).text = "本科"
            document.add_paragraph("项目经历")
            document.save(path)
            text = extract_text_from_file(path, "docx")
        self.assertLess(text.index("教育经历"), text.index("示例大学"))
        self.assertLess(text.index("本科"), text.index("项目经历"))


class ArchiveAPITests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client_context = TestClient(app)
        cls.client = cls.client_context.__enter__()

    @classmethod
    def tearDownClass(cls):
        cls.client_context.__exit__(None, None, None)

    def workspace(self):
        demo = self.client.post("/api/auth/demo").json()
        result = self.client.post("/api/auth/workspace", headers={"X-Auth-Token": demo["token"]})
        self.assertEqual(result.status_code, 200)
        return result.json()

    def test_both_system_examples_are_complete_in_archive_and_original(self):
        demo = self.client.post("/api/auth/demo").json()
        headers = {"X-Auth-Token": demo["token"]}
        rows = self.client.get("/api/resumes", headers=headers).json()["items"]
        for resume_id in ("cv_demo_linzhixia", "cv_sample"):
            with self.subTest(resume_id=resume_id):
                item = next(row for row in rows if row["id"] == resume_id)
                sections = item["profile"]["resume_sections"]
                self.assertTrue(all(sections.values()))
                source = self.client.get(f"/api/resumes/{resume_id}/source", headers=headers).json()["raw_text"]
                self.assertIn("AI 知识检索助手原型", source)
                self.assertIn("虚构", source)
                self.assertNotIn("（示例简历，无真实文件）", source)

    def test_upload_persistence_and_access_isolation(self):
        first, other = self.workspace(), self.workspace()
        headers = {"X-Auth-Token": first["token"]}
        upload = self.client.post("/api/resumes/upload", headers=headers, files={"file": ("resume.txt", RESUME.encode(), "text/plain")})
        self.assertEqual(upload.status_code, 200, upload.text)
        data = upload.json()
        self.assertTrue(data["profile"]["resume_sections"]["education"])
        self.assertFalse(data["profile"]["parse_meta"]["used"])
        self.assertNotIn("raw_text", data)
        resumed = self.client.post("/api/auth/demo", headers=headers).json()
        self.assertEqual(resumed["user"]["id"], first["user"]["id"])
        self.assertEqual(resumed["token"], first["token"])
        listing = self.client.get("/api/resumes", headers=headers).json()
        self.assertIn(data["id"], [row["id"] for row in listing["items"]])
        source = self.client.get(f'/api/resumes/{data["id"]}/source', headers=headers)
        self.assertEqual(source.json()["raw_text"], RESUME)
        unauthorized = self.client.get(f'/api/resumes/{data["id"]}/source', headers={"X-Auth-Token": other["token"]})
        self.assertEqual(unauthorized.status_code, 404)
        other_listing = self.client.get("/api/resumes", headers={"X-Auth-Token": other["token"]}).json()
        self.assertNotIn(data["id"], [row["id"] for row in other_listing["items"]])
        shared = self.client.post("/api/auth/demo").json()
        demo_listing = self.client.get("/api/resumes", headers={"X-Auth-Token": shared["token"]}).json()
        self.assertNotIn(data["id"], [row["id"] for row in demo_listing["items"]])
        for route, body in (("profile", {"profile": {"skills": ["虚构能力"]}}), ("append-story", {"section": "新经历", "text": "未写在文件中的经历"})):
            response = (self.client.put if route == "profile" else self.client.post)(f'/api/resumes/{data["id"]}/{route}', headers=headers, json=body)
            self.assertEqual(response.status_code, 409)

    def test_bad_file_is_not_reported_as_parsed(self):
        workspace = self.workspace()
        headers = {"X-Auth-Token": workspace["token"]}
        response = self.client.post("/api/resumes/upload", headers=headers, files={"file": ("empty.txt", b"   \n", "text/plain")})
        self.assertEqual(response.status_code, 400)
        rows = self.client.get("/api/resumes", headers=headers).json()["items"]
        self.assertEqual(next(row for row in rows if not row["is_sample"])["status"], "failed")
        response = self.client.post("/api/resumes/upload", headers=headers, files={"file": ("old.doc", b"not-a-doc", "application/msword")})
        self.assertEqual(response.status_code, 400)

    def test_demo_matching_uses_the_same_facts_as_archive(self):
        demo = self.client.post("/api/auth/demo").json()
        headers = {"X-Auth-Token": demo["token"]}
        rows = self.client.get("/api/resumes", headers=headers).json()["items"]
        archive = next(row for row in rows if row["id"].startswith("cv_demo_"))
        with patch("server.routers.matches._run_match_core", return_value={"ok": True}) as run:
            response = self.client.post("/api/match/run", headers=headers, json={
                "resume_id": archive["id"], "jd_text": demo["presets"][0]["jd_text"], "auto_report": False,
            })
        self.assertEqual(response.status_code, 200, response.text)
        # Legacy prebuilt results must not bypass the new evidence policy.
        run.assert_called_once()
        profile = run.call_args.kwargs["resume_profile"]
        self.assertEqual(profile["resume_sections"], archive["profile"]["resume_sections"])
        self.assertEqual(profile["personality_traits"], [])
        self.assertEqual(profile["suitable_roles"], [])

    def test_interview_receives_validated_archive_facts(self):
        demo = self.client.post("/api/auth/demo").json()
        headers = {"X-Auth-Token": demo["token"]}
        rows = self.client.get("/api/resumes", headers=headers).json()["items"]
        archive = next(row for row in rows if row["id"].startswith("cv_demo_"))
        with patch("server.llm.llm.complete_json", return_value=({}, {"used": False})) as complete:
            response = self.client.post(f'/api/resumes/{archive["id"]}/interview-report', headers=headers, json={
                "transcript": "我负责客户门户的需求梳理与灰度上线。", "questions": [],
            })
        self.assertEqual(response.status_code, 200)
        complete.assert_not_called()
        self.assertTrue(response.json()["llm"]["local"])


if __name__ == "__main__":
    unittest.main()
